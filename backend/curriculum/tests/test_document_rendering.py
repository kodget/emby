"""Real document bytes and local converters; no database or provider calls."""
from pathlib import Path
import tempfile
import unittest
import os
import subprocess
from django.test import SimpleTestCase
from unittest.mock import patch
import fitz
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
from curriculum.services.document_rendering import render_document, cleanup_render, office_executable
from curriculum.services.converter import SlideConverter
from curriculum.services.slide_conversion_pipeline import SlideConversionPipeline


def make_document(directory, kind):
    source = Path(directory) / f'lecture.{kind}'
    if kind == 'ppt':
        modern = make_document(directory, 'pptx')
        with tempfile.TemporaryDirectory(prefix='emby-office-test-') as profile:
            subprocess.run([office_executable(), f'-env:UserInstallation={Path(profile).resolve().as_uri()}',
                            '--headless', '--convert-to', 'ppt:MS PowerPoint 97', '--outdir', str(directory), str(modern)],
                           capture_output=True, check=True, timeout=120,
                           creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
        if not source.read_bytes().startswith(bytes.fromhex('d0cf11e0')):
            raise AssertionError('Legacy fixture was not a real binary PowerPoint file')
        return source
    if kind == 'pdf':
        with fitz.open() as document:
            for number in range(1, 4):
                page = document.new_page(width=400, height=300)
                page.insert_text((30, 40), f'Reader fixture page {number}')
                page.draw_rect(fitz.Rect(40, 70, 150, 130), color=(1, 0, 0), fill=(1, 0, 0))
            document.save(source)
    elif kind == 'docx':
        document = Document()
        for number in range(1, 4):
            document.add_paragraph(f'Reader fixture page {number}')
            if number < 3:
                document.add_page_break()
        document.save(source)
    else:
        document = Presentation()
        for number in range(1, 4):
            slide = document.slides.add_slide(document.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = f'Reader fixture page {number}'
        document.save(source)
    return source


class DocumentRenderingTests(SimpleTestCase):
    def test_scanned_uploads_render_without_starting_ocr_in_either_pipeline(self):
        from curriculum.tests.test_document_extraction import pdf_bytes
        with tempfile.TemporaryDirectory(prefix='emby-fixture-') as directory:
            source = Path(directory) / 'scanned.pdf'
            source.write_bytes(pdf_bytes(scanned=(1, 2), count=2))
            with patch.object(fitz.Page, 'get_textpage_ocr') as ocr:
                course = SlideConversionPipeline.process_slide(str(source), 'fixture', 'pdf')
                success, standalone = SlideConverter.convert_document(str(source), 'pdf', 'fixture')
            try:
                self.assertTrue(course['success'])
                self.assertTrue(success)
                self.assertEqual(course['text_extraction']['missing_pages'], [1, 2])
                self.assertEqual(standalone['page_count'], 2)
                self.assertTrue(all(Path(p).is_file() for p in course['image_paths']))
                ocr.assert_not_called()
            finally:
                cleanup_render(course)
                cleanup_render(standalone)

    def test_real_pdf_word_and_powerpoint_keep_every_page_in_order(self):
        for kind in ('pdf', 'docx', 'pptx', 'ppt'):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory(prefix='emby-fixture-') as directory:
                source = make_document(directory, kind)
                original = source.read_bytes()
                result = render_document(source, kind, Path(directory) / 'rendered')
                self.assertEqual(result['page_count'], 3)
                for number, (path, text) in enumerate(zip(result['image_paths'], result['page_texts']), 1):
                    self.assertIn(f'Reader fixture page {number}', text)
                    with Image.open(path) as image:
                        self.assertGreater(image.width, 600)
                        self.assertGreater(len(image.convert('L').getcolors(256) or []), 1)
                self.assertEqual(source.read_bytes(), original)

    def test_standalone_images_survive_until_the_caller_finishes_storing_them(self):
        with tempfile.TemporaryDirectory(prefix='emby-fixture-') as directory:
            source = make_document(directory, 'pdf')
            success, result = SlideConverter.convert_document(str(source), 'pdf', 'fixture')
            try:
                self.assertTrue(success, result)
                self.assertEqual(result['page_count'], 3)
                self.assertTrue(all(Path(p).is_file() for p in result['slide_images']))
                root = Path(result['temp_dir'])
            finally:
                cleanup_render(result)
            self.assertFalse(root.exists())
            self.assertTrue(source.is_file())

    def test_worker_accepts_office_document_instead_of_rejecting_non_pdf(self):
        with tempfile.TemporaryDirectory(prefix='emby-fixture-') as directory:
            source = make_document(directory, 'pptx')
            with patch('curriculum.services.slide_conversion_pipeline._download_from_cloudinary', return_value=source.read_bytes()):
                result = SlideConversionPipeline.process_slide('https://storage.invalid/source', 'fixture', 'pptx')
            try:
                self.assertTrue(result['success'], result)
                self.assertEqual(result['page_count'], 3)
                self.assertIn('Reader fixture page 3', result['text_content'])
                self.assertTrue(all(Path(p).is_file() for p in result['image_paths']))
            finally:
                cleanup_render(result)

    def test_corrupt_document_fails_without_returning_deleted_or_partial_pages(self):
        with tempfile.TemporaryDirectory(prefix='emby-fixture-') as directory:
            source = Path(directory) / 'broken.pdf'
            source.write_bytes(b'not a PDF')
            success, result = SlideConverter.convert_document(str(source), 'pdf', 'broken')
            self.assertFalse(success)
            self.assertIsNone(result['temp_dir'])
            self.assertEqual(result['slide_images'], [])

    def test_missing_office_converter_has_an_explicit_failure(self):
        with tempfile.TemporaryDirectory(prefix='emby-fixture-') as directory:
            source = make_document(directory, 'pptx')
            with patch('curriculum.services.document_rendering.office_executable', side_effect=RuntimeError('unavailable')):
                success, result = SlideConverter.convert_document(str(source), 'pptx', 'unavailable')
            self.assertFalse(success)
            self.assertEqual(result['page_count'], 0)
            self.assertIsNone(result['temp_dir'])

    def test_cleanup_cannot_remove_an_unrelated_directory(self):
        with tempfile.TemporaryDirectory(prefix='unrelated-fixture-') as directory:
            with self.assertRaises(ValueError):
                cleanup_render({'temp_dir': directory})
            self.assertTrue(Path(directory).is_dir())
