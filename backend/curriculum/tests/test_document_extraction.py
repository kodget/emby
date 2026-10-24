"""Real page content, scanned images and failure coverage, without external services."""
import os
from pathlib import Path
import tempfile
from unittest.mock import patch

import fitz
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from django.test import SimpleTestCase, override_settings

from curriculum.content_extractor import extract_text_from_slide, get_slide_full_text
from curriculum.services.document_extraction import (
    ExtractionError, content_text, extract_document, is_complete, ocr_options,
)
from curriculum.services.document_rendering import render_document
from curriculum.tests.test_document_rendering import make_document


def pdf_bytes(*, scanned=(), count=5):
    with fitz.open() as document:
        for number in range(1, count + 1):
            page = document.new_page(width=600, height=600)
            text = f'Anatomy lecture page {number}\nThe heart pumps blood around the body.\nFinal marker number {number}.'
            if number in scanned:
                with fitz.open() as scan:
                    scan_page = scan.new_page(width=600, height=600)
                    scan_page.insert_text((40, 60), text, fontsize=20)
                    image = scan_page.get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes('png')
                page.insert_image(page.rect, stream=image)
            else:
                page.insert_text((40, 60), text, fontsize=20)
        return document.tobytes()


class DocumentExtractionTests(SimpleTestCase):
    def test_native_extraction_has_no_fifteen_thousand_character_cutoff(self):
        with fitz.open() as doc:
            for number in range(30):
                page = doc.new_page()
                page.insert_text((30, 40), ('Reliable extraction preserves the whole lecture.\n' * 25) + f'LAST MARKER {number}', fontsize=10)
            raw = doc.tobytes()
        text = get_slide_full_text(raw, 'application/pdf')
        self.assertGreater(len(text), 15000)
        self.assertIn('LAST MARKER 29', text)
        with self.assertRaises(ExtractionError) as caught:
            get_slide_full_text(raw, 'pdf', max_chars=15000)
        self.assertEqual(caught.exception.code, 'text_limit')

    def test_scanned_and_mixed_documents_include_pages_after_the_third(self):
        # This is a real Tesseract run; missing language data must fail the test.
        ocr_options()
        for scanned in ((1, 2, 3, 4, 5), (2, 4, 5)):
            with self.subTest(scanned=scanned):
                result = extract_document(pdf_bytes(scanned=scanned), 'pdf')
                self.assertTrue(is_complete(result), result['text_extraction'])
                self.assertEqual(result['text_extraction']['ocr_pages'], list(scanned))
                for number, page in enumerate(result['pages'], 1):
                    text = ' '.join(page['text'].split())
                    self.assertIn(f'Anatomy lecture page {number}', text)
                    self.assertIn(f'Final marker number {number}', text)

    def test_mixed_native_header_and_scanned_body_both_survive(self):
        with fitz.open(stream=pdf_bytes(scanned=(1,), count=1), filetype='pdf') as doc:
            doc[0].insert_text((30, 25), 'Native heading')
            result = extract_document(doc.tobytes(), 'pdf')
        self.assertTrue(is_complete(result))
        self.assertIn('Native heading', result['pages'][0]['text'])
        self.assertIn('The heart pumps blood', result['pages'][0]['text'])

    def test_ocr_unavailable_reports_exact_missing_pages_and_preserves_native_text(self):
        with patch.dict(os.environ, {'TESSDATA_PREFIX': '/no-such-language-data'}):
            result = extract_document(pdf_bytes(scanned=(2, 4)), 'pdf')
        self.assertFalse(is_complete(result))
        self.assertEqual(result['text_extraction']['missing_pages'], [2, 4])
        self.assertEqual(result['text_extraction']['status'], 'partial')
        self.assertEqual(result['text_extraction']['error_counts'], {'ocr_unavailable': 2})
        self.assertEqual(result['pages'][1]['error_code'], 'ocr_unavailable')
        self.assertIn('Final marker number 5', result['pages'][4]['text'])
        with self.assertRaises(ExtractionError):
            content_text(result, require_complete=True)

    def test_debug_mode_discovers_repository_local_language_data(self):
        tessdata = Path(__file__).resolve().parents[3] / '.test-tessdata'
        trained_data = tessdata / 'eng.traineddata'
        def is_file(path):
            return path == trained_data

        with patch.dict(os.environ, {'TESSDATA_PREFIX': ''}), patch.object(Path, 'is_file', is_file):
            with override_settings(DEBUG=True):
                self.assertEqual(ocr_options()['tessdata'], str(tessdata))
            with override_settings(DEBUG=False), self.assertRaises(ExtractionError) as caught:
                ocr_options()
        self.assertEqual(caught.exception.code, 'ocr_unavailable')

    def test_one_ocr_failure_does_not_skip_later_pages(self):
        original = fitz.Page.get_textpage_ocr
        def failing(page, **kwargs):
            if page.number == 1:
                raise RuntimeError('private provider diagnostic')
            return original(page, **kwargs)
        with patch.object(fitz.Page, 'get_textpage_ocr', failing):
            result = extract_document(pdf_bytes(scanned=(2, 5)), 'pdf')
        self.assertEqual(result['text_extraction']['missing_pages'], [2])
        self.assertIn('Final marker number 5', result['pages'][4]['text'])
        self.assertNotIn('private provider', str(result))

    def test_native_extraction_failure_is_a_page_failure_not_a_blank_page(self):
        with patch.object(fitz.Page, 'get_text', side_effect=RuntimeError('bad font')):
            result = extract_document(pdf_bytes(count=1), 'pdf')
        self.assertEqual(result['pages'][0]['error_code'], 'native_failed')
        self.assertEqual(result['text_extraction']['status'], 'failed')

    def test_true_blank_page_is_preserved_and_nonblank_unreadable_page_is_flagged(self):
        with fitz.open() as doc:
            doc.new_page()
            page = doc.new_page()
            page.draw_rect(fitz.Rect(10, 10, 200, 200), color=(0, 0, 0), fill=(0, 0, 0))
            result = extract_document(doc.tobytes(), 'pdf')
        self.assertEqual(result['pages'][0]['extraction_status'], 'blank')
        self.assertEqual(result['pages'][1]['error_code'], 'unreadable')
        self.assertEqual(result['text_extraction']['missing_pages'], [2])

    def test_corrupt_encrypted_and_unsupported_sources_fail_explicitly(self):
        with fitz.open() as doc:
            doc.new_page()
            encrypted = doc.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='password')
        for source, kind, code in ((b'broken', 'pdf', 'invalid_document'),
                                    (encrypted, 'pdf', 'encrypted'), (b'text', 'txt', 'unsupported')):
            with self.subTest(code=code), self.assertRaises(ExtractionError) as caught:
                extract_document(source, kind)
            self.assertEqual(caught.exception.code, code)

    @patch('requests.get')
    def test_http_error_is_not_parsed_as_a_document(self, download):
        download.return_value.status_code = 403
        download.return_value.content = b'private error page'
        with self.assertRaises(ExtractionError) as caught:
            extract_text_from_slide('https://storage.invalid/source?signature=secret', 'pdf')
        self.assertEqual(caught.exception.code, 'source_unavailable')
        self.assertNotIn('secret', str(caught.exception))

    def test_office_page_numbers_tables_and_grouped_shapes_match_rendering(self):
        with tempfile.TemporaryDirectory(prefix='emby-extraction-fixture-') as directory:
            word = Document()
            word.add_paragraph('Word page one')
            word.add_table(rows=1, cols=1).cell(0, 0).text = 'Table cell information'
            word.add_page_break()
            word.add_paragraph('Word final page')
            path = Path(directory) / 'table.docx'
            word.save(path)
            result = extract_document(path, 'docx')
            self.assertEqual(len(result['pages']), 2)
            self.assertIn('Table cell information', result['pages'][0]['text'])
            self.assertIn('Word final page', result['pages'][1]['text'])
            for kind in ('pptx', 'ppt'):
                source = make_document(directory, kind)
                pages = extract_text_from_slide(source, kind)
                self.assertEqual(len(pages), 3)
                self.assertIn('Reader fixture page 3', pages[-1]['text'])
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            group = slide.shapes.add_group_shape()
            group.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = 'Grouped text survives'
            slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(5), Inches(1)).table.cell(0, 0).text = 'PowerPoint table survives'
            path = Path(directory) / 'grouped.pptx'
            deck.save(path)
            text = get_slide_full_text(path, 'pptx')
            self.assertIn('Grouped text survives', text)
            self.assertIn('PowerPoint table survives', text)

    def test_images_remain_readable_when_ocr_is_unavailable(self):
        with tempfile.TemporaryDirectory(prefix='emby-extraction-fixture-') as directory:
            path = Path(directory) / 'scanned.pdf'
            path.write_bytes(pdf_bytes(scanned=(1, 2), count=2))
            with patch.dict(os.environ, {'TESSDATA_PREFIX': '/missing-language-data'}):
                result = render_document(path, 'pdf', Path(directory) / 'pages', ocr=True)
            self.assertEqual(len(result['image_paths']), 2)
            self.assertTrue(all(Path(image).is_file() for image in result['image_paths']))
            self.assertEqual(result['text_extraction']['missing_pages'], [1, 2])

    def test_legacy_or_incomplete_manifest_cannot_claim_whole_document_coverage(self):
        result = extract_document(pdf_bytes(), 'pdf')
        self.assertTrue(is_complete(result))
        result['pages'].pop()
        self.assertFalse(is_complete(result))
        self.assertFalse(is_complete({'text': 'first three pages only'}))
